from __future__ import annotations

import shutil
import sys
import unittest
import uuid
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
TEST_TEMP_ROOT = PROJECT_ROOT / "temp_sync" / "unit-tests"
TEST_TEMP_ROOT.mkdir(parents=True, exist_ok=True)
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

from excel_reader import snapshot_excel  # noqa: E402
from sync_excel import (  # noqa: E402
    E1_FILE,
    E1_HEADERS,
    E1_SHEET,
    build_enter1,
    check_lock_files,
    enter1_extra_headers,
    resolve_sync_settings,
)
from validators import Issues  # noqa: E402
from image_extractor import ImageAnchor  # noqa: E402


@contextmanager
def test_directory():
    root = TEST_TEMP_ROOT / uuid.uuid4().hex
    root.mkdir(parents=False, exist_ok=False)
    try:
        yield root
    finally:
        shutil.rmtree(root, ignore_errors=True)


class SnapshotSafetyTests(unittest.TestCase):
    def test_enter1_extra_headers_follow_excel_column_order(self) -> None:
        headers = {
            "国家": 1,
            "产品类别": 2,
            "其他特殊标签": 19,
            "新增文字": 20,
            "新增图片": 22,
        }
        self.assertEqual(enter1_extra_headers(headers), [("新增文字", 20), ("新增图片", 22)])

    def test_enter1_duplicate_extra_header_is_rejected(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = E1_SHEET
        for column, header in enumerate(E1_HEADERS + ["重复字段", "重复字段"], start=1):
            worksheet.cell(2, column, header)
        issues = Issues()
        with patch("sync_excel.open_excel", return_value=workbook):
            records, _, _ = build_enter1(Path("unused.xlsx"), PROJECT_ROOT, issues)
        self.assertEqual(records, [])
        self.assertTrue(any(issue.reason == "Duplicate header is not allowed" for issue in issues.errors))

    def test_enter1_extra_column_becomes_ordered_json_field(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = E1_SHEET
        for column, header in enumerate(E1_HEADERS + ["新增文字", "新增图片"], start=1):
            worksheet.cell(2, column, header)
        fixed_values = [
            "中国", "电动工具", "测试产品", "Corded", "", "TEST-1", "123", "220V",
            "MM/YYYY", "", "", "", "", "", "", "", "IP20", "", "",
        ]
        for column, value in enumerate(fixed_values + ["测试内容", ""], start=1):
            worksheet.cell(3, column, value)

        issues = Issues()
        with patch("sync_excel.open_excel", return_value=workbook), patch(
            "sync_excel.extract_image_anchors", return_value=[]
        ):
            records, _, _ = build_enter1(Path("unused.xlsx"), PROJECT_ROOT, issues)

        self.assertFalse(issues.errors)
        self.assertEqual([field["column"] for field in records[0]["extraFields"]], [20, 21])
        self.assertEqual(records[0]["extraFields"][0]["label"]["zh"], "新增文字")
        self.assertEqual(records[0]["extraFields"][0]["value"]["zh"], "测试内容")
        self.assertEqual(records[0]["extraFields"][1]["images"], [])

    def test_enter1_extra_column_accepts_anchored_images(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = E1_SHEET
        for column, header in enumerate(E1_HEADERS + ["新增图片区"], start=1):
            worksheet.cell(2, column, header)
        values = [
            "中国", "电动工具", "测试产品", "Corded", "", "TEST-2", "456", "220V",
            "MM/YYYY", "", "", "", "", "", "", "", "IP20", "", "", "",
        ]
        for column, value in enumerate(values, start=1):
            worksheet.cell(3, column, value)
        anchor = ImageAnchor(E1_SHEET, 3, 20, 0, 0, 1, "xl/media/image1.png", "abc", b"image")

        issues = Issues()
        with patch("sync_excel.open_excel", return_value=workbook), patch(
            "sync_excel.extract_image_anchors", return_value=[anchor]
        ), patch("sync_excel.save_as_png"):
            records, _, generated = build_enter1(Path("unused.xlsx"), PROJECT_ROOT, issues)

        self.assertFalse(issues.errors)
        self.assertEqual(len(records[0]["extraFields"][0]["images"]), 1)
        self.assertEqual(records[0]["extraFields"][0]["images"][0], generated[0].as_posix())
        self.assertIn("_extra_c20_01.png", generated[0].as_posix())

    def test_source_change_during_snapshot_is_rejected(self) -> None:
        with test_directory() as root:
            source = root / "source.xlsx"
            destination = root / "snapshot.xlsx"
            source.write_bytes(b"original workbook bytes")
            real_copy2 = shutil.copy2

            def copy_then_change(source_path: Path, destination_path: Path):
                result = real_copy2(source_path, destination_path)
                with Path(source_path).open("ab") as handle:
                    handle.write(b"changed while snapshotting")
                return result

            with patch("excel_reader.shutil.copy2", side_effect=copy_then_change):
                with self.assertRaisesRegex(RuntimeError, "Source changed while being copied"):
                    snapshot_excel(source, destination)

    def test_excel_lock_file_blocks_sync(self) -> None:
        with test_directory() as root:
            (root / ("~$" + E1_FILE)).touch()
            with self.assertRaisesRegex(RuntimeError, "Excel lock file detected"):
                check_lock_files(root)

    def test_production_mode_rejects_project_test_excel(self) -> None:
        with self.assertRaisesRegex(RuntimeError, "Production mode cannot publish from ./excel"):
            resolve_sync_settings(PROJECT_ROOT, PROJECT_ROOT / "excel", True)


if __name__ == "__main__":
    unittest.main()
