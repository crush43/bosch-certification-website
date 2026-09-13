"""Excel text and table reading helpers."""

from __future__ import annotations

import hashlib
import re
import shutil
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


def clean_text(value: Any) -> str:
    """Convert a cell value, including rich text, to a cleaned string."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        text = value.isoformat()
    else:
        text = str(value)
    text = text.strip()
    return "" if text == "/" else text


def localized_zh(value: Any) -> dict[str, str]:
    return {"zh": clean_text(value), "en": "", "de": ""}


def normalize_id_part(value: Any) -> str:
    text = unicodedata.normalize("NFKC", clean_text(value)).casefold()
    return re.sub(r"\s+", " ", text).strip()


def stable_hash(parts: Iterable[Any], length: int = 12) -> str:
    normalized = "\x1f".join(normalize_id_part(part) for part in parts)
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()[:length]


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def file_modified_at(path: Path) -> str:
    return datetime.fromtimestamp(path.stat().st_mtime).astimezone().isoformat(timespec="seconds")


def snapshot_excel(source: Path, destination: Path) -> dict[str, Any]:
    """Copy one source workbook and prove it did not change during the copy."""
    before = source.stat()
    shutil.copy2(source, destination)
    after = source.stat()
    if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
        raise RuntimeError(f"Source changed while being copied: {source.name}")
    source_hash = file_sha256(source)
    copied_hash = file_sha256(destination)
    final = source.stat()
    if (after.st_size, after.st_mtime_ns) != (final.st_size, final.st_mtime_ns):
        raise RuntimeError(f"Source changed while being verified: {source.name}")
    if source_hash != copied_hash:
        raise RuntimeError(f"Snapshot checksum mismatch: {source.name}")
    return {
        "file": source.name,
        "sourcePath": str(source.resolve()),
        "size": final.st_size,
        "modifiedAt": file_modified_at(source),
        "snapshotAt": datetime.now().astimezone().isoformat(timespec="seconds"),
        "sha256": source_hash,
    }


def open_excel(path: Path):
    # rich_text=True preserves all runs; clean_text(str(value)) produces display text.
    return load_workbook(path, data_only=False, read_only=False, rich_text=True)


def header_map(worksheet, header_row: int) -> dict[str, int]:
    result: dict[str, int] = {}
    for cell in worksheet[header_row]:
        name = clean_text(cell.value)
        if name:
            result[name] = cell.column
    return result


def row_is_blank(worksheet, row_number: int, columns: Iterable[int]) -> bool:
    return all(not clean_text(worksheet.cell(row_number, column).value) for column in columns)
